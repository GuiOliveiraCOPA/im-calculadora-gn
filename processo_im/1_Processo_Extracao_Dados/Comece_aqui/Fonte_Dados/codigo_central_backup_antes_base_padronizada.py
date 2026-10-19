from __future__ import annotations
import argparse
import os
import re
import subprocess
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
TODAY = datetime.now().strftime("%Y-%m-%d")
OUTPUT_DIR = ROOT_DIR / "PlanilhasUF"
DEFAULT_OUTPUT = ROOT_DIR / "PlanilhaResultado" / "Planilha_GN_Atualizada.xlsx"
CONTROL_SHEET = "controle"
PROTECTION_PASSWORD = "GN"

def execution_metadata() -> Dict[str, str]:
    user = os.environ.get("NOME_RESPONSAVEL") or os.environ.get("GN_EXEC_USER") or os.environ.get("USERNAME") or ""
    computer = os.environ.get("GN_EXEC_COMPUTER") or os.environ.get("COMPUTERNAME") or ""
    started_at = os.environ.get("GN_EXEC_STARTED_AT") or ""
    return {
        "usuario": user,
        "usuario_windows": os.environ.get("GN_EXEC_USER") or os.environ.get("USERNAME") or "",
        "computador": computer,
        "inicio_execucao": started_at,
    }


@dataclass(frozen=True)
class Extractor:
    name: str
    script: str
    output_glob: str
    states: tuple[str, ...]
    sheet_map: Optional[Dict[str, str]] = None


EXTRACTORS: tuple[Extractor, ...] = (
    Extractor("SP", "tabela_sp.py", "Sp_*.xlsx", ("SP",)),
    Extractor("RJ", "tabela_rj.py", "Rj_*.xlsx", ("RJ",)),
    Extractor("ES", "tabela_esgas.py", "Esgas_*.xlsx", ("ES",)),
    Extractor("MG", "tabela_gasmig.py", "Gasmig_*.xlsx", ("MG",)),
    Extractor("PR", "tabela_compagas.py", "Compagas_*.xlsx", ("PR",)),
    Extractor("SC", "tabela_scgas.py", "Scgas_*.xlsx", ("SC",)),
    Extractor("RS", "tabela_sulgas.py", "Sulgas_*.xlsx", ("RS",)),
    Extractor("MS", "tabela_msgas.py", "Msgas_*.xlsx", ("MS",)),
    Extractor("MT", "tabela_mtgas.py", "Mtgas_*.xlsx", ("MT",)),
    Extractor("AM", "tabela_cigas.py", "Cigas_*.xlsx", ("AM",)),
    Extractor("AL", "tabela_algas.py", "Algas_*.xlsx", ("AL",)),
    Extractor("BA", "tabela_bahiagas.py", "Bahiagas_*.xlsx", ("BA",)),
    Extractor("CE", "tabela_cegas.py", "Cegas_*.xlsx", ("CE",)),
    Extractor("PB", "tabela_pbgas.py", "Pbgas_*.xlsx", ("PB",)),
    Extractor("PE", "tabela_copergas.py", "Copergas_*.xlsx", ("PE",)),
    Extractor("RN", "tabela_potigas.py", "Potigas_*.xlsx", ("RN",)),
    Extractor("SE", "tabela_sergipegas.py", "Sergipegas_*.xlsx", ("SE",)),
)

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Roda os extratores estaduais e consolida tudo em um Excel central.")
    parser.add_argument(
        "--only",
        help="Lista de UFs para rodar, separadas por virgula. Exemplo: --only SP,RJ,ES",
    )
    parser.add_argument(
        "--skip-run",
        action="store_true",
        help="Nao executa os scripts; usa os Excels mais recentes ja existentes na pasta.",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        default=900,
        help="Tempo maximo, em segundos, para cada extrator.",
    )
    parser.add_argument(
        "--output",
        default=str(DEFAULT_OUTPUT),
        help="Caminho do Excel central gerado.",
    )
    return parser.parse_args()


def selected_extractors(only: Optional[str]) -> List[Extractor]:
    if not only:
        return list(EXTRACTORS)
    wanted = {item.strip().upper() for item in only.split(",") if item.strip()}
    return [extractor for extractor in EXTRACTORS if wanted.intersection(extractor.states)]


def latest_file(pattern: str) -> Optional[Path]:
    files = list(BASE_DIR.glob(pattern))
    if OUTPUT_DIR.exists():
        files.extend(OUTPUT_DIR.glob(pattern))
    if not files:
        return None
    return max(files, key=lambda path: path.stat().st_mtime)


def next_available_path(path: Path) -> Path:
    if not path.exists():
        return path

    timestamp = datetime.now().strftime("%H%M%S")
    for index in range(1, 100):
        suffix = timestamp if index == 1 else f"{timestamp}_{index}"
        candidate = path.with_name(f"{path.stem}_{suffix}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Nao consegui criar nome unico para {path.name}")


def move_workbook_to_output_folder(source_path: Path) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    source_path = source_path.resolve()
    destination = (OUTPUT_DIR / source_path.name).resolve()

    if source_path == destination:
        return source_path

    try:
        if destination.exists():
            destination.unlink()
    except OSError:
        destination = next_available_path(destination)

    try:
        shutil.move(str(source_path), str(destination))
    except OSError:
        destination = next_available_path(destination)
        shutil.move(str(source_path), str(destination))
    return destination

def run_extractor(extractor: Extractor, timeout: int, skip_run: bool) -> dict:
    script_path = BASE_DIR / extractor.script
    started_at = time.time()
    result = {
        "extractor": extractor.name,
        "script": extractor.script,
        "states": ",".join(extractor.states),
        "status": "pendente",
        "arquivo": "",
        "mensagem": "",
    }

    if not script_path.exists():
        result["status"] = "script ausente"
        result["mensagem"] = f"Arquivo nao encontrado: {extractor.script}"
        return result

    if not skip_run:
        env = os.environ.copy()
        env["PYTHONIOENCODING"] = "utf-8"
        completed = subprocess.run(
            [sys.executable, str(script_path)],
            cwd=BASE_DIR,
            env=env,
            text=True,
            encoding="utf-8",
            errors="replace",
            capture_output=True,
            timeout=timeout,
        )
        if completed.returncode != 0:
            result["status"] = "erro ao executar"
            stdout = (completed.stdout or "").strip()
            stderr = (completed.stderr or "").strip()
            result["mensagem"] = (stderr or stdout)[-1000:]
            fallback = latest_file(extractor.output_glob)
            if fallback:
                result["arquivo"] = str(fallback)
            return result

    output_file = latest_file(extractor.output_glob)
    if not output_file:
        result["status"] = "sem excel"
        result["mensagem"] = f"Nenhum arquivo encontrado com padrao {extractor.output_glob}"
        return result

    if not skip_run and output_file.stat().st_mtime < started_at - 2:
        result["status"] = "excel antigo"
        result["mensagem"] = "Usei o Excel mais recente encontrado, mas ele parece anterior a esta execucao."
    else:
        result["status"] = "ok"

    result["arquivo"] = str(output_file)
    return result


def ensure_sheet(workbook: Workbook, title: str) -> Worksheet:
    safe_title = title[:31]
    if safe_title in workbook.sheetnames:
        return workbook[safe_title]
    return workbook.create_sheet(safe_title)


def is_sheet_empty(sheet: Worksheet) -> bool:
    return sheet.max_row == 1 and sheet.max_column == 1 and sheet.cell(1, 1).value is None


def append_values(sheet: Worksheet, values: list) -> None:
    if is_sheet_empty(sheet):
        for column_index, value in enumerate(values, start=1):
            sheet.cell(1, column_index).value = value
        return
    sheet.append(values)

def append_blank_row(sheet: Worksheet) -> None:
    if sheet.max_row > 1 or sheet.cell(1, 1).value is not None:
        sheet.append([])


def numbered_sheet_title(workbook: Workbook, base_title: str, number: int) -> str:
    attempt = max(number, 1)
    while True:
        suffix = "" if attempt == 1 else str(attempt)
        available = 31 - len(suffix)
        candidate = f"{base_title[:available]}{suffix}"
        if candidate not in workbook.sheetnames:
            return candidate
        attempt += 1


def create_numbered_sheet(workbook: Workbook, base_title: str, number: int) -> Worksheet:
    return workbook.create_sheet(numbered_sheet_title(workbook, base_title, number))


def has_daily_unit(value: object) -> bool:
    text = str(value or "").lower()
    return any(token in text for token in [
        "m3_dia", "m³/dia", "m3/dia", "mł/dia", "m³ / dia", "m3 / dia", "mł / dia", "/dia"
    ])


def replace_daily_unit_with_monthly(value: object) -> object:
    if value is None:
        return value
    text = str(value)
    replacements = [
        ("m3_dia", "m3_mes"),
        ("M3_DIA", "M3_MES"),
        ("m³/dia", "m³/mês"),
        ("M³/DIA", "m³/mês"),
        ("m3/dia", "m³/mês"),
        ("M3/DIA", "m³/mês"),
        ("mł/dia", "m³/mês"),
        ("MŁ/DIA", "m³/mês"),
        ("m³ / dia", "m³/mês"),
        ("M³ / DIA", "m³/mês"),
        ("m3 / dia", "m³/mês"),
        ("M3 / DIA", "m³/mês"),
        ("mł / dia", "m³/mês"),
        ("MŁ / DIA", "m³/mês"),
        ("/dia", "/mês"),
        ("/DIA", "/mês"),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
    return text


def parse_br_number_token(token: str) -> tuple[float, int]:
    token = token.strip()
    decimal_places = 0
    if "," in token:
        decimal_places = len(token.rsplit(",", 1)[1])
        normalized = token.replace(".", "").replace(",", ".")
    elif "." in token:
        normalized = token.replace(".", "")
    else:
        normalized = token
    return float(normalized), decimal_places


def format_br_number(value: float, decimal_places: int) -> str:
    formatted = f"{value:,.{decimal_places}f}"
    return formatted.replace(",", "X").replace(".", ",").replace("X", ".")


def convert_number_token_to_monthly(match) -> str:
    token = match.group(0)
    try:
        value, decimal_places = parse_br_number_token(token)
    except ValueError:
        return token
    return format_br_number(value * 30, decimal_places)


def convert_daily_text_to_monthly(value: object, convert_plain_number: bool = False) -> object:
    if value is None:
        return value
    text = str(value)
    if not text.strip():
        return value

    range_like = bool(re.search(r"\b(a|até|ate|acima|de)\b", text.lower())) or has_daily_unit(text)
    if not convert_plain_number and not range_like:
        return replace_daily_unit_with_monthly(text)

    converted = re.sub(r"\d[\d\.,]*", convert_number_token_to_monthly, text)
    return replace_daily_unit_with_monthly(converted)


def is_specific_daily_range_header(value: object) -> bool:
    text = str(value or "").lower()
    if not has_daily_unit(text):
        return False
    return any(token in text for token in ["inicial", "final", "min", "max", "mínimo", "máximo", "minimo", "maximo"])


def row_has_daily_unit(values: list) -> bool:
    return any(has_daily_unit(value) for value in values)


def convert_daily_row_to_monthly(values: list, daily_numeric_columns: set[int], daily_table: bool) -> tuple[list, set[int], bool]:
    original = list(values)
    converted = list(values)
    row_daily = row_has_daily_unit(original)

    if row_daily:
        daily_table = True
        for index, value in enumerate(original):
            if is_specific_daily_range_header(value):
                daily_numeric_columns.add(index)
                converted[index] = replace_daily_unit_with_monthly(value)
            else:
                converted[index] = convert_daily_text_to_monthly(value, convert_plain_number=False)
        return converted, daily_numeric_columns, daily_table

    if daily_table:
        for index, value in enumerate(converted):
            if index in daily_numeric_columns:
                converted[index] = convert_daily_text_to_monthly(value, convert_plain_number=True)
            elif value is not None:
                text = str(value)
                if re.search(r"\b(a|até|ate|acima|de)\b", text.lower()):
                    converted[index] = convert_daily_text_to_monthly(value, convert_plain_number=False)

    return converted, daily_numeric_columns, daily_table

def same_row(left: list, right: list) -> bool:
    max_len = max(len(left), len(right))
    for index in range(max_len):
        left_value = "" if index >= len(left) or left[index] is None else str(left[index]).strip()
        right_value = "" if index >= len(right) or right[index] is None else str(right[index]).strip()
        if left_value != right_value:
            return False
    return True


def is_blank_row_values(values: list) -> bool:
    return not any(value is not None and str(value).strip() != "" for value in values)


def is_title_like_row(values: list) -> bool:
    non_empty = [value for value in values if value is not None and str(value).strip() != ""]
    return len(non_empty) == 1


def normalize_source_rows(rows: List[tuple]) -> List[list]:
    normalized: List[list] = []
    materialized = [list(row) for row in rows]

    for index, values in enumerate(materialized):
        if is_blank_row_values(values):
            if not normalized or is_blank_row_values(normalized[-1]):
                continue

            next_non_blank = None
            for next_values in materialized[index + 1:]:
                if not is_blank_row_values(next_values):
                    next_non_blank = next_values
                    break

            if is_title_like_row(normalized[-1]) and next_non_blank is not None:
                continue

            normalized.append([])
            continue

        normalized.append(values)

    while normalized and is_blank_row_values(normalized[-1]):
        normalized.pop()

    return normalized


def append_source_sheet(
    target: Worksheet,
    source: Worksheet,
    source_file: Path,
    source_title: str,
    include_heading: bool,
    skip_duplicate_header: bool = False,
) -> int:
    appended = 0
    if include_heading:
        append_blank_row(target)
        append_values(target, [source_title])
        appended += 1

    rows = normalize_source_rows(list(source.iter_rows(values_only=True)))
    if skip_duplicate_header and rows and target.max_row >= 1:
        source_header = list(rows[0])
        target_header = [target.cell(1, column).value for column in range(1, target.max_column + 1)]
        if target.cell(1, 1).value is not None and same_row(source_header, target_header):
            rows = rows[1:]

    daily_numeric_columns: set[int] = set()
    daily_table = False

    for values in rows:
        if is_blank_row_values(values):
            append_blank_row(target)
            daily_numeric_columns = set()
            daily_table = False
            appended += 1
            continue
        values, daily_numeric_columns, daily_table = convert_daily_row_to_monthly(values, daily_numeric_columns, daily_table)
        append_values(target, values)
        appended += 1
    return appended

def copy_workbook_to_master(master: Workbook, extractor: Extractor, source_path: Path, control_rows: List[dict]) -> None:
    source_wb = load_workbook(source_path, data_only=False)

    if extractor.sheet_map:
        for source_sheet, target_sheet in extractor.sheet_map.items():
            if source_sheet not in source_wb.sheetnames:
                control_rows.append({
                    "extractor": extractor.name,
                    "estado": target_sheet,
                    "arquivo": str(source_path),
                    "status": "aba ausente",
                    "detalhe": f"Aba {source_sheet} nao encontrada no arquivo de origem.",
                })
                continue
            target = ensure_sheet(master, target_sheet)
            rows = append_source_sheet(target, source_wb[source_sheet], source_path, source_sheet, include_heading=False, skip_duplicate_header=True)
            control_rows.append({
                "extractor": extractor.name,
                "estado": target_sheet,
                "arquivo": str(source_path),
                "status": "copiado",
                "detalhe": f"{rows} linhas copiadas da aba {source_sheet}.",
            })
        return

    target_state = extractor.states[0]
    target = ensure_sheet(master, target_state)
    usable_sheets = [name for name in source_wb.sheetnames if name.lower() != CONTROL_SHEET]
    include_heading = len(usable_sheets) > 1

    for sheet_name in usable_sheets:
        rows = append_source_sheet(target, source_wb[sheet_name], source_path, sheet_name, include_heading=include_heading)
        control_rows.append({
            "extractor": extractor.name,
            "estado": target_state,
            "arquivo": str(source_path),
            "status": "copiado",
            "detalhe": f"{rows} linhas copiadas da aba {sheet_name}.",
        })


def write_control_sheet(master: Workbook, run_rows: List[dict], copy_rows: List[dict]) -> None:
    sheet = ensure_sheet(master, CONTROL_SHEET)
    sheet.append(["tipo", "extractor", "estado", "script", "arquivo", "status", "detalhe"])

    metadata = execution_metadata()
    sheet.append(["metadado", "usuario_windows", "", "", "", metadata["usuario"], ""])
    sheet.append(["metadado", "computador", "", "", "", metadata["computador"], ""])
    sheet.append(["metadado", "inicio_execucao", "", "", "", metadata["inicio_execucao"], ""])

    for row in run_rows:
        sheet.append([
            "execucao",
            row.get("extractor", ""),
            row.get("states", ""),
            row.get("script", ""),
            row.get("arquivo", ""),
            row.get("status", ""),
            row.get("mensagem", ""),
        ])

    for row in copy_rows:
        sheet.append([
            "consolidacao",
            row.get("extractor", ""),
            row.get("estado", ""),
            "",
            row.get("arquivo", ""),
            row.get("status", ""),
            row.get("detalhe", ""),
        ])


def autosize(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        for column_index, column_cells in enumerate(sheet.columns, start=1):
            letter = get_column_letter(column_index)
            width = 10
            for cell in column_cells:
                if cell.value is not None:
                    width = max(width, len(str(cell.value)) + 2)
            sheet.column_dimensions[letter].width = min(width, 60)


def apply_standard_format(workbook: Workbook) -> None:
    title_fill = PatternFill("solid", fgColor="E2F0D9")
    header_fill = PatternFill("solid", fgColor="EDEDED")
    title_font = Font(bold=True, size=12, color="375623")
    header_font = Font(bold=True)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            values = [cell.value for cell in row]
            non_empty = [value for value in values if value is not None and str(value).strip() != ""]
            if not non_empty:
                continue

            if len(non_empty) == 1:
                for cell in row:
                    cell.font = title_font
                    cell.fill = title_fill
                    cell.alignment = Alignment(vertical="center")
                continue

            first_value = str(row[0].value or "").strip().lower()
            looks_like_header = (
                "segmento" in first_value
                or "faixa" in first_value
                or first_value == "tipo"
                or "mínimo" in first_value
                or "minimo" in first_value
                or "descri" in first_value
            )
            if looks_like_header:
                for cell in row:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

def remove_default_sheet(workbook: Workbook) -> None:
    if "Sheet" in workbook.sheetnames and len(workbook.sheetnames) > 1:
        del workbook["Sheet"]



def create_warning_sheet(workbook: Workbook) -> None:
    if "Aviso" in workbook.sheetnames:
        del workbook["Aviso"]

    sheet = workbook.create_sheet("Aviso", 0)
    sheet.sheet_view.showGridLines = False
    metadata = execution_metadata()

    title_fill = PatternFill("solid", fgColor="E2F0D9")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    note_fill = PatternFill("solid", fgColor="F8FBF3")
    title_font = Font(bold=True, size=24, color="375623")
    subtitle_font = Font(bold=True, size=13, color="7F6000")
    body_font = Font(size=12, color="1F1F1F")
    small_font = Font(size=10, italic=True, color="666666")

    sheet.merge_cells("A1:H1")
    sheet["A1"] = "Atenção!"
    sheet["A1"].font = title_font
    sheet["A1"].fill = title_fill
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 36

    sheet.merge_cells("A3:H4")
    sheet["A3"] = "Esta planilha está vinculada a outra base de dados. Altere somente se tiver certeza do impacto."
    sheet["A3"].font = subtitle_font
    sheet["A3"].fill = warning_fill
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    messages = [
        "Evite editar, apagar ou reorganizar abas, colunas, cabeçalhos e linhas da planilha final.",
        "Para atualizar os dados, use sempre o arquivo AtualizarDados.bat.",
        "Se precisar fazer análises manuais, trabalhe em uma cópia do arquivo.",
        "Em caso de erro na atualização, consulte a pasta Execucoes_CMD e envie o arquivo mais recente para suporte.",
    ]

    for index, message in enumerate(messages, start=6):
        cell = sheet.cell(row=index, column=1)
        cell.value = message
        cell.font = body_font
        cell.fill = note_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=8)
        sheet.row_dimensions[index].height = 24

    sheet.merge_cells("A12:H12")
    sheet["A12"] = f"Arquivo atualizado em: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    sheet["A12"].font = small_font
    sheet["A12"].alignment = Alignment(horizontal="center")

    audit_lines = [
        f"Atualizado por: {metadata['usuario'] or 'usuário não identificado'}",
        f"Computador: {metadata['computador'] or 'computador não identificado'}",
        f"Início da execução: {metadata['inicio_execucao'] or 'não informado'}",
    ]
    for offset, value in enumerate(audit_lines, start=13):
        sheet.merge_cells(start_row=offset, start_column=1, end_row=offset, end_column=8)
        sheet.cell(row=offset, column=1).value = value
        sheet.cell(row=offset, column=1).font = small_font
        sheet.cell(row=offset, column=1).alignment = Alignment(horizontal="center")

    for column in range(1, 9):
        sheet.column_dimensions[chr(64 + column)].width = 18

    workbook.active = workbook.sheetnames.index("Aviso")

def protect_workbook_for_viewing(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.protection.sheet = True
        sheet.protection.set_password(PROTECTION_PASSWORD)

def save_master_workbook(workbook: Workbook, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
        return output_path
    except PermissionError as exc:
        raise PermissionError(
            f"Nao consegui atualizar {output_path.name}. Feche esse arquivo no Excel e execute novamente."
        ) from exc

def main() -> None:
    args = parse_args()
    chosen = selected_extractors(args.only)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = ROOT_DIR / output_path

    print(f"Extratores selecionados: {', '.join(extractor.name for extractor in chosen)}")
    print(f"Pasta das planilhas individuais: {OUTPUT_DIR}")

    master = Workbook()
    run_rows: List[dict] = []
    copy_rows: List[dict] = []

    for extractor in chosen:
        print(f"\nRodando {extractor.name}...")
        run_info = run_extractor(extractor, timeout=args.timeout, skip_run=args.skip_run)
        run_rows.append(run_info)
        print(f"Status: {run_info['status']}")

        source_file = run_info.get("arquivo")
        if not source_file or not Path(source_file).exists():
            continue
        try:
            organized_file = move_workbook_to_output_folder(Path(source_file))
            run_info["arquivo"] = str(organized_file)
            source_file = str(organized_file)
        except Exception as exc:
            previous_message = run_info.get("mensagem", "")
            run_info["mensagem"] = f"{previous_message} | Nao consegui mover para PlanilhasUF: {exc}".strip(" |")
        try:
            copy_workbook_to_master(master, extractor, Path(source_file), copy_rows)
        except Exception as exc:
            copy_rows.append({
                "extractor": extractor.name,
                "estado": ",".join(extractor.states),
                "arquivo": source_file,
                "status": "erro ao consolidar",
                "detalhe": str(exc),
            })

    write_control_sheet(master, run_rows, copy_rows)
    remove_default_sheet(master)
    apply_standard_format(master)
    create_warning_sheet(master)
    autosize(master)
    protect_workbook_for_viewing(master)
    output_path = save_master_workbook(master, output_path)

    print("\nConcluido.")
    print(f"Arquivo central gerado em: {output_path}")


if __name__ == "__main__":
    main()


















