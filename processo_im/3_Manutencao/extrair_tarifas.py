from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE_FILE = PROJECT_ROOT / "2_Bases_Origem" / "Tarifario_GN.xlsx"
DEFAULT_OUTPUT_FILE = PROJECT_ROOT / "5_Calculadora_Final" / "tarifas.js"


def as_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def as_text(value):
    return "" if value is None else str(value).strip()


def as_date_text(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return as_text(value)


def get_index(headers):
    index = {as_text(name): position for position, name in enumerate(headers)}
    if "Concessionaria" not in index and "ss" in index:
        index["Concessionaria"] = index["ss"]
    return index


def main() -> None:
    source_file = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE_FILE
    output_file = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_OUTPUT_FILE

    if not source_file.exists():
        raise FileNotFoundError(f"Base não encontrada: {source_file}")

    workbook = load_workbook(source_file, data_only=True, read_only=True)
    if "Tarifario" not in workbook.sheetnames:
        raise ValueError("A planilha precisa ter uma aba chamada 'Tarifario'.")

    sheet = workbook["Tarifario"]
    rows = sheet.iter_rows(values_only=True)
    headers = [as_text(item) for item in next(rows)]
    index = get_index(headers)

    required_columns = [
        "Segmento Copa",
        "Concessionaria",
        "Preço da Concessionaria é:",
        "Classe",
        "UF",
        "Concessionária",
        "Segmento",
        "Faixa Mín (m3)",
        "Faixa Máx (m3)",
        "Tarifa (R$/m³)",
        "Termo Fixo (R$)",
        "dt_delib",
        "PIS/Cofins",
        "ICMS",
        "Ordem",
        "Tarifa em cascata?",
        "Fonte",
    ]
    missing = [column for column in required_columns if column not in index]
    if missing:
        raise ValueError(f"Colunas ausentes na base: {missing}")

    records = []
    for row in rows:
        if not any(row):
            continue
        record = {
            "segmento_copa": as_text(row[index["Segmento Copa"]]),
            "concessionaria_label": as_text(row[index["Concessionaria"]]),
            "preco_concessionaria": as_text(row[index["Preço da Concessionaria é:"]]),
            "classe": as_number(row[index["Classe"]]),
            "uf": as_text(row[index["UF"]]),
            "concessionaria": as_text(row[index["Concessionária"]]),
            "segmento": as_text(row[index["Segmento"]]),
            "faixa_min": as_number(row[index["Faixa Mín (m3)"]]),
            "faixa_max": as_number(row[index["Faixa Máx (m3)"]]),
            "tarifa": as_number(row[index["Tarifa (R$/m³)"]]),
            "termo_fixo": as_number(row[index["Termo Fixo (R$)"]]) or 0,
            "dt_delib": as_date_text(row[index["dt_delib"]]),
            "pis_cofins": as_number(row[index["PIS/Cofins"]]) or 0,
            "icms": as_number(row[index["ICMS"]]) or 0,
            "ordem": as_number(row[index["Ordem"]]) or 0,
            "cascata": as_text(row[index["Tarifa em cascata?"]]).lower().startswith("sim"),
            "fonte": as_text(row[index["Fonte"]]),
        }
        if record["uf"] and record["concessionaria"] and record["faixa_min"] is not None:
            records.append(record)

    payload = {
        "source": str(source_file),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "records": records,
    }
    output_file.write_text(
        "window.TARIFAS_GN = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    print(f"Gerado: {output_file}")
    print(f"Registros: {len(records)}")


if __name__ == "__main__":
    main()


