from __future__ import annotations

import json
import re
import subprocess
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
PROCESS_DIR = PROJECT_ROOT / "1_Processo_Extracao_Dados" / "Comece_aqui"
FINAL_EXTRACT_FILE = PROCESS_DIR / "Planilha_Final" / "Planilha_GN_Atualizada.xlsx"
SOURCE_TARIFF_FILE = PROJECT_ROOT / "2_Bases_Origem" / "Tarifario_GN.xlsx"
OUTPUT_TARIFFS_JS = PROJECT_ROOT / "5_Calculadora_Final" / "tarifas.js"
OUTPUT_TARIFFS_JSON = PROJECT_ROOT / "6_Calculadora_Flask" / "data" / "tarifas.json"
CENTRAL_CODE_DIR = PROCESS_DIR / "Fonte_Dados"
MAX_RANGE_VALUE = 99999999999.0


def normalize(value) -> str:
    text = "" if value is None else str(value)
    replacements = {
        "M?n": "Min",
        "M?x": "Max",
        "m?n": "min",
        "m?x": "max",
        "mł": "m3",
        "m³": "m3",
        "mÂ³": "m3",
        "Â³": "3",
        "Mediçăo": "Medicao",
        "mediçăo": "medicao",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = text.lower().replace("r$", "rs")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def as_text(value):
    return "" if value is None else str(value).strip()


def as_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("R$", "").replace("%", "")
    if not text or text in {"-", "–", "—"}:
        return None
    text = text.replace(".", "").replace(",", ".") if "," in text else text
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def is_close(a, b, tolerance=0.02):
    if a is None or b is None:
        return False
    return abs(float(a) - float(b)) <= tolerance


def range_key(min_value, max_value):
    min_value = 0 if min_value is None else float(min_value)
    max_value = MAX_RANGE_VALUE if max_value is None else float(max_value)
    return (round(min_value, 2), round(max_value, 2))


def import_central_module():
    sys.path.insert(0, str(CENTRAL_CODE_DIR))
    import codigo_central as central  # type: ignore

    central.normalize_key = normalize
    return central


def extract_records_from_final_workbook():
    if not FINAL_EXTRACT_FILE.exists():
        raise FileNotFoundError(f"Planilha final do .bat não encontrada: {FINAL_EXTRACT_FILE}")

    central = import_central_module()
    workbook = load_workbook(FINAL_EXTRACT_FILE, data_only=True, read_only=True)
    records = []
    for sheet_name in workbook.sheetnames:
        if sheet_name in {"Aviso", "controle"}:
            continue
        records.extend(central.standardize_worksheet(sheet_name, workbook[sheet_name]))

    liquid_candidates = []
    for record in records:
        tariff_key = normalize(record.get("Tipo_Tarifa"))
        if any(skip in tariff_key for skip in ("com_impostos", "c_tributos", "tributos", "bruta")) and not any(
            keep in tariff_key for keep in ("sem_impostos", "s_tributos", "ex_tributos", "sem_imposto")
        ):
            continue
        if any(keep in tariff_key for keep in ("sem_impostos", "s_tributos", "ex_tributos", "sem_imposto", "total_sem_impostos", "tarifa", "valor_variavel", "tarifa_rsm3")):
            liquid_candidates.append(record)
    return liquid_candidates


def read_existing_tariff_rows():
    if not SOURCE_TARIFF_FILE.exists():
        raise FileNotFoundError(f"Base consolidada não encontrada: {SOURCE_TARIFF_FILE}")

    workbook = load_workbook(SOURCE_TARIFF_FILE, data_only=True)
    sheet = workbook["Tarifario"]
    headers = [as_text(cell.value) for cell in sheet[1]]
    index = {header: pos + 1 for pos, header in enumerate(headers)}
    required = [
        "Segmento Copa", "Concessionaria", "Preço da Concessionaria é:", "Classe", "UF", "Concessionária",
        "Segmento", "Faixa Mín (m3)", "Faixa Máx (m3)", "Tarifa (R$/m³)", "Termo Fixo (R$)",
        "dt_delib", "PIS/Cofins", "ICMS", "Ordem", "Tarifa em cascata?", "Fonte"
    ]
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"Colunas ausentes em Tarifario_GN.xlsx: {missing}")
    return workbook, sheet, index


def candidate_score(existing, candidate):
    if normalize(existing["uf"]) != normalize(candidate.get("UF")):
        return -1

    score = 0
    if normalize(existing["segmento_copa"]) and normalize(existing["segmento_copa"]) == normalize(candidate.get("Segmento_Copa")):
        score += 4

    existing_segment = normalize(existing["segmento"])
    candidate_segment = normalize(candidate.get("Segmento"))
    if existing_segment and candidate_segment:
        if existing_segment == candidate_segment:
            score += 3
        elif existing_segment in candidate_segment or candidate_segment in existing_segment:
            score += 2

    existing_label = normalize(existing["concessionaria_label"])
    candidate_label = normalize(candidate.get("Concessionaria"))
    if existing_label and candidate_label:
        if existing_label == candidate_label:
            score += 3
        elif existing_label in candidate_label or candidate_label in existing_label:
            score += 2

    if is_close(existing["faixa_min"], candidate.get("Faixa_Min_m3")):
        score += 3
    if is_close(existing["faixa_max"], candidate.get("Faixa_Max_m3"), tolerance=1.0):
        score += 3

    return score


def build_updated_records():
    candidates = extract_records_from_final_workbook()
    workbook, sheet, index = read_existing_tariff_rows()

    updated_count = 0
    unmatched_count = 0
    records = []
    update_details = []

    for row_number in range(2, sheet.max_row + 1):
        existing = {
            "segmento_copa": as_text(sheet.cell(row_number, index["Segmento Copa"]).value),
            "concessionaria_label": as_text(sheet.cell(row_number, index["Concessionaria"]).value),
            "preco_concessionaria": as_text(sheet.cell(row_number, index["Preço da Concessionaria é:"]).value),
            "classe": as_number(sheet.cell(row_number, index["Classe"]).value),
            "uf": as_text(sheet.cell(row_number, index["UF"]).value),
            "concessionaria": as_text(sheet.cell(row_number, index["Concessionária"]).value),
            "segmento": as_text(sheet.cell(row_number, index["Segmento"]).value),
            "faixa_min": as_number(sheet.cell(row_number, index["Faixa Mín (m3)"]).value),
            "faixa_max": as_number(sheet.cell(row_number, index["Faixa Máx (m3)"]).value),
            "tarifa": as_number(sheet.cell(row_number, index["Tarifa (R$/m³)"]).value),
            "termo_fixo": as_number(sheet.cell(row_number, index["Termo Fixo (R$)"]).value) or 0,
            "dt_delib": as_text(sheet.cell(row_number, index["dt_delib"]).value),
            "pis_cofins": as_number(sheet.cell(row_number, index["PIS/Cofins"]).value) or 0,
            "icms": as_number(sheet.cell(row_number, index["ICMS"]).value) or 0,
            "ordem": as_number(sheet.cell(row_number, index["Ordem"]).value) or 0,
            "cascata": as_text(sheet.cell(row_number, index["Tarifa em cascata?"]).value).lower().startswith("sim"),
            "fonte": as_text(sheet.cell(row_number, index["Fonte"]).value),
        }

        if not (existing["uf"] and existing["concessionaria"] and existing["faixa_min"] is not None and existing["tarifa"] is not None):
            continue

        scored = [(candidate_score(existing, candidate), candidate) for candidate in candidates]
        scored = [(score, candidate) for score, candidate in scored if score >= 9]
        scored.sort(key=lambda item: item[0], reverse=True)

        if scored:
            best_score, best = scored[0]
            new_tariff = best.get("Tarifa")
            new_fixed = best.get("Termo_Fixo")
            if new_tariff is not None:
                if existing["tarifa"] != float(new_tariff):
                    update_details.append(
                        f"{existing['uf']} | {existing['segmento_copa']} | {existing['concessionaria_label']} | faixa {existing['faixa_min']}-{existing['faixa_max']}: tarifa {existing['tarifa']} -> {new_tariff}"
                    )
                existing["tarifa"] = float(new_tariff)
                existing["termo_fixo"] = float(new_fixed or 0)
                updated_count += 1
        else:
            unmatched_count += 1

        records.append(existing)

    return records, candidates, updated_count, unmatched_count, update_details


def write_tariffs_js(records):
    payload = {
        "source": str(FINAL_EXTRACT_FILE),
        "schema_reference": str(SOURCE_TARIFF_FILE),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "records": records,
    }
    OUTPUT_TARIFFS_JS.write_text(
        "window.TARIFAS_GN = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    # Mesma base para a versão Flask (6_Calculadora_Flask). Grava em arquivo
    # temporário e troca no final, para o servidor nunca ler um JSON pela metade.
    OUTPUT_TARIFFS_JSON.parent.mkdir(parents=True, exist_ok=True)
    temporario = OUTPUT_TARIFFS_JSON.with_suffix(".json.tmp")
    temporario.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporario.replace(OUTPUT_TARIFFS_JSON)


def main():
    run_bat = "--run-bat" in sys.argv
    if run_bat:
        subprocess.run(["cmd", "/c", str(PROCESS_DIR / "AtualizarDados.bat")], cwd=PROCESS_DIR, check=True)

    records, candidates, updated_count, unmatched_count, update_details = build_updated_records()
    write_tariffs_js(records)

    print(f"Planilha final lida: {FINAL_EXTRACT_FILE}")
    print(f"Candidatos extraídos da planilha final: {len(candidates)}")
    print(f"Registros finais no tarifas.js: {len(records)}")
    print(f"Registros atualizados por casamento automático: {updated_count}")
    print(f"Registros mantidos pela base de referência: {unmatched_count}")
    print(f"Gerado: {OUTPUT_TARIFFS_JS}")
    print(f"Gerado: {OUTPUT_TARIFFS_JSON}")
    if update_details:
        print("Exemplos de alterações:")
        for detail in update_details[:10]:
            print(f"- {detail}")


if __name__ == "__main__":
    main()



